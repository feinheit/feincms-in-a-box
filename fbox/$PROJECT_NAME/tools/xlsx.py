"""
Example usage inside Django's admin panel::

    from django.contrib import admin
    from django.utils.translation import ugettext_lazy as _

    from ${PROJECT_NAME} import models
    from ${PROJECT_NAME}.tools.xlsx import XLSDocument


    class AttendanceAdmin(admin.ModelAdmin):
        list_filter = ('event',)
        actions = ('export_selected', 'delete_selected')

        def export_selected(self, request, queryset):
            xlsx = XLSDocument()
            xlsx.table_from_queryset(queryset)
            return xlsx.to_response('attendance.xlsx')
        export_selected.short_description = _('export selected')


    admin.site.register(models.Event)
    admin.site.register(models.Attendance, AttendanceAdmin)
"""

from __future__ import absolute_import, unicode_literals

from datetime import date
from decimal import Decimal
from io import BytesIO
from openpyxl import Workbook

from django.http import HttpResponse


class XLSDocument(object):
    def __init__(self):
        self.workbook = Workbook(optimized_write=True)
        self.sheet = None

    def add_sheet(self, title=None):
        self.sheet = self.workbook.create_sheet(title=title)

    def table(self, titles, rows):
        if titles:
            self.sheet.append(titles)

        for row in rows:
            processed = []
            for i, value in enumerate(row):
                if isinstance(value, date):
                    processed.append(value.strftime('%Y-%m-%d'))
                elif isinstance(value, (int, float, Decimal)):
                    processed.append(value)
                elif value is None:
                    processed.append('-')
                else:
                    processed.append(unicode(value).strip())

            self.sheet.append(processed)

    def table_from_queryset(self, queryset):
        opts = queryset.model._meta

        titles = ['str']
        titles.extend(field.name for field in opts.fields)

        data = []
        for instance in queryset:
            row = ['%s' % instance]
            for field in opts.fields:
                if field.choices:
                    row.append(
                        getattr(instance, 'get_%s_display' % field.name)())
                else:
                    row.append(getattr(instance, field.name))

            data.append(row)

        self.add_sheet('%s' % opts.verbose_name_plural)
        self.table(titles, data)

    def to_response(self, filename):
        output = BytesIO()
        self.workbook.save(output)
        response = HttpResponse(
            output.getvalue(),
            content_type=(
                'application/vnd.openxmlformats-officedocument.'
                'spreadsheetml.sheet'),
        )
        output.close()
        response['Content-Disposition'] = 'attachment; filename="%s"' % (
            filename,
        )
        return response
